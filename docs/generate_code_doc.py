"""
将 HappyEnglish 全部源代码整理成 Word 文档
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import glob

BASE_DIR = r'D:\claude code\HappyEnglish'
OUTPUT_FILE = r'D:\claude code\HappyEnglish\HappyEnglish源代码.docx'

# 需要整理的核心文件
FILES = [
    ('index.html', '主页面'),
    ('src/index.html', 'src/主页面'),
    ('src/app.js', 'src/应用入口'),
    ('src/js/app.js', 'src/js/应用'),
    ('test.html', '测试页面'),
    ('batch-test.html', '批量测试'),
    ('profile.html', '个人中心'),
    ('review.html', '复习页面'),
    ('words.html', '词汇页面'),
    ('src/pages/home.html', '页面/首页'),
    ('src/pages/etymology.html', '页面/词根词缀'),
    ('src/pages/profile.html', '页面/个人中心'),
    ('src/pages/report.html', '页面/报告'),
    ('src/pages/review.html', '页面/复习'),
    ('src/pages/story.html', '页面/故事'),
    ('src/pages/achievements.html', '页面/成就'),
    ('src/js/ai-service.js', 'js/AI服务'),
    ('src/js/etymology.js', 'js/词根词缀'),
    ('src/js/db.js', 'js/数据库'),
    ('src/js/router.js', 'js/路由'),
    ('src/js/store.js', 'js/状态管理'),
    ('src/js/level.js', 'js/等级'),
    ('src/js/spaced-repetition.js', 'js/间隔复习'),
    ('src/js/analytics.js', 'js/分析'),
    ('src/js/learning-companion.js', 'js/学习伙伴'),
    ('src/js/trigger-engine.js', 'js/触发引擎'),
    ('src/js/memory-profile.js', 'js/记忆画像'),
    ('src/js/personalized-suggestions.js', 'js/个性化建议'),
    ('src/js/achievement-system.js', 'js/成就系统'),
    ('src/components/companion-widget.js', '组件/学习伙伴'),
    ('src/data/vocabulary_full.js', '数据/词汇库'),
    ('src/data/etymology.js', '数据/词根词缀'),
    ('src/data/grammar.js', '数据/语法'),
    ('src/services/speech.js', '服务/语音'),
]

def create_code_document():
    doc = Document()

    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('HappyEnglish 源代码')
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x66, 0x7E, 0xEA)

    # 副标题
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('智能英语学习助手 - 完整源代码')
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x4A, 0x55, 0x69)

    doc.add_paragraph()

    # 目录
    toc_title = doc.add_paragraph()
    run = toc_title.add_run('文件目录')
    run.bold = True
    run.font.size = Pt(16)

    for i, (filepath, desc) in enumerate(FILES, 1):
        full_path = os.path.join(BASE_DIR, filepath)
        exists = '✓' if os.path.exists(full_path) else '✗'
        p = doc.add_paragraph()
        p.add_run(f'{i:2d}. {desc} ({filepath}) {exists}')

    doc.add_page_break()

    # 逐个文件添加内容
    for filepath, desc in FILES:
        full_path = os.path.join(BASE_DIR, filepath)

        # 添加文件标题
        heading = doc.add_paragraph()
        run = heading.add_run(f'=== {filepath} ===')
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x66, 0x7E, 0xEA)

        if not os.path.exists(full_path):
            p = doc.add_paragraph()
            p.add_run(f'[文件不存在: {filepath}]')
            doc.add_paragraph()
            continue

        try:
            with open(full_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # 添加代码内容
            code_para = doc.add_paragraph()
            code_para.style = 'Normal'

            # 设置代码样式 - 使用文本块保留格式
            lines = content.split('\n')
            for line in lines:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(line)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        except Exception as e:
            p = doc.add_paragraph()
            p.add_run(f'[读取失败: {str(e)}]')

        doc.add_paragraph()
        doc.add_paragraph()

    doc.save(OUTPUT_FILE)
    print(f'源代码文档已生成: {OUTPUT_FILE}')

if __name__ == '__main__':
    create_code_document()